"""Create editable PowerPoint and PDF competition materials from versioned sources."""
import json
import re
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts"
INK, CARD, CYAN, WHITE, MUTED = "0A1020", "121B2E", "31BDF2", "EDF1FA", "ADB9CE"


def text(slide, value, x, y, width, height, size=24, color=WHITE, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    paragraph = tf.paragraphs[0]
    paragraph.text = value
    paragraph.font.name = "DejaVu Sans"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    return box


def shape(slide, x, y, width, height, color, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(width), Inches(height))
    item.fill.solid()
    item.fill.fore_color.rgb = RGBColor.from_string(color)
    item.line.fill.background()
    return item


def build_deck():
    content = json.loads((ROOT / "presentation/deck.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, entry in enumerate(content["slides"]):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(INK)
        shape(slide, 11.95, -0.4, 2.1, 2.1, CARD, MSO_SHAPE.OVAL)
        text(slide, entry["kicker"], .65, .5, 11.4, .3, 12, CYAN, True)
        text(slide, entry["title"], .65, 1.08, 11.2, 1.55, 34 if len(entry["title"]) < 63 else 30, WHITE, True)
        if i == 8:
            metrics = [("PRECISION", "87,5%"), ("RECALL", "100%"), ("F1-SCORE", "0,933"), ("FPR", "14,3%")]
            for j, (label, value) in enumerate(metrics):
                x = .65 + j * 3.08
                shape(slide, x, 3.0, 2.83, 1.62, CARD)
                text(slide, label, x+.2, 3.2, 2.4, .3, 12, MUTED)
                text(slide, value, x+.2, 3.65, 2.4, .65, 33, CYAN, True)
            text(slide, "Synthetic test · 42 строки / 14 шаблонов · TP 21 / FP 3 / TN 18 / FN 0", .65, 4.95, 11.5, .7, 18, WHITE)
            text(slide, "Правила сильнее ML. Нужен новый независимый слепой тест.", .65, 5.75, 11.5, .65, 19, CYAN)
        elif i == 5:
            boxes = [("Frontend", .65), ("FastAPI", 3.75), ("Rules + ML", 6.85), ("Risk Engine", 9.95)]
            for label, x in boxes:
                shape(slide, x, 3.1, 2.55, 1.0, CARD)
                text(slide, label, x+.2, 3.38, 2.2, .4, 20, WHITE, True)
                if x < 9: text(slide, "→", x+2.65, 3.32, .4, .5, 25, CYAN)
            text(slide, "Результат → PostgreSQL → история / dashboard / PDF", .65, 4.65, 11.5, .7, 23, CYAN)
            text(slide, "Исходные сообщения не сохраняются. Ссылки не открываются.", .65, 5.55, 11.5, .8, 20, MUTED)
        else:
            for j, point in enumerate(entry["points"]):
                y = 2.95 + j * .78
                text(slide, f"{j+1:02d}", .65, y, .5, .4, 15, CYAN, True)
                text(slide, point, 1.35, y-.03, 10.85, .72, 21 if len(point) < 67 else 19, WHITE)
        text(slide, "QALQAN FINANCE SECURITY  /  RESEARCH PROTOTYPE", .65, 7.02, 10.7, .25, 9, MUTED)
        text(slide, f"{i+1:02d} / 12", 11.8, 7.02, .9, .25, 10, CYAN)
        slide.notes_slide.notes_text_frame.text = entry["notes"]
    prs.save(OUTPUT / "Qalqan-pitch-deck.pptx")


def build_pdf(source):
    font = ROOT / "backend/fonts/QalqanSans.ttf"
    pdfmetrics.registerFont(TTFont("Qalqan", str(font)))
    styles = {
        "title": ParagraphStyle("title", fontName="Qalqan", fontSize=20, leading=26, textColor=colors.HexColor("#" + INK)),
        # A darker cyan remains legible on the white printable page.
        "subtitle": ParagraphStyle("subtitle", fontName="Qalqan", fontSize=11, leading=16, textColor=colors.HexColor("#087BA5")),
        "body": ParagraphStyle("body", fontName="Qalqan", fontSize=9.5, leading=14, textColor=colors.HexColor("#" + CARD)),
    }
    story = []
    for paragraph in source.read_text(encoding="utf-8").split("\n\n"):
        if not paragraph.strip():
            continue
        if paragraph.startswith("## "):
            style, value = "subtitle", paragraph[3:]
        elif paragraph.startswith("# "):
            style, value = "title", paragraph[2:]
        else:
            style, value = "body", paragraph.replace("\n", " ")
        value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
        story.append(Paragraph(escape(value), styles[style]))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(OUTPUT / (source.stem + ".pdf")), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=32, bottomMargin=30, title=source.stem).build(story)


if __name__ == "__main__":
    OUTPUT.mkdir(exist_ok=True)
    build_deck()
    for name in ("summary-ru.md", "summary-en.md", "one-pager.md"):
        build_pdf(ROOT / "presentation" / name)
    print(f"Built PowerPoint and three PDFs in {OUTPUT}")
